import pandas as pd
import pytest
from openpyxl import load_workbook

from src.run import main
from src.stock_template import write_stock_template

HEADERS = ['Product', 'Stock On Hand', 'Count Date', 'Weeks To Arrive',
           'Min Order Qty', 'On Order']
AS_OF = pd.Timestamp('2026-10-01')


def _sales(rows):
    frame = pd.DataFrame(rows, columns=['product', 'date', 'quantity'])
    return frame.assign(date=pd.to_datetime(frame['date']))


def _rows(path):
    ws = load_workbook(path)['Stock & Orders']
    return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]


def test_template_has_the_stock_sheet_with_the_expected_headers(tmp_path):
    sales = _sales([('A', '2026-09-01', 1)])
    path = write_stock_template(sales, tmp_path / 't.xlsx', AS_OF)
    ws = load_workbook(path)['Stock & Orders']
    assert [c.value for c in ws[1]] == HEADERS


def test_template_lists_products_sold_in_the_last_year_alphabetically(tmp_path):
    sales = _sales([
        ('Zeta', '2026-09-01', 1), ('Alpha', '2026-03-01', 2),
        ('Retired', '2024-01-01', 1),
    ])
    path = write_stock_template(sales, tmp_path / 't.xlsx', AS_OF)
    assert [row[0] for row in _rows(path)] == ['Alpha', 'Zeta']


def test_refund_only_products_are_not_listed(tmp_path):
    sales = _sales([('Alpha', '2026-09-01', 1), ('Refunded', '2026-09-02', -1)])
    path = write_stock_template(sales, tmp_path / 't.xlsx', AS_OF)
    assert [row[0] for row in _rows(path)] == ['Alpha']


def test_only_the_minimum_order_and_on_order_are_prefilled(tmp_path):
    sales = _sales([('Alpha', '2026-09-01', 1)])
    path = write_stock_template(sales, tmp_path / 't.xlsx', AS_OF)
    assert _rows(path) == [['Alpha', None, None, None, 1, 0]]


def test_template_includes_a_how_to_sheet(tmp_path):
    sales = _sales([('Alpha', '2026-09-01', 1)])
    path = write_stock_template(sales, tmp_path / 't.xlsx', AS_OF)
    assert len(load_workbook(path).sheetnames) == 2


def test_command_writes_the_template_without_building_a_report(
        workbook_path, tmp_path):
    target = tmp_path / 'stock_template.xlsx'
    demo = tmp_path / 'demo'
    code = main([
        '--profile', 'synthetic', '--file', str(workbook_path),
        '--stock-template', str(target), '--demo-output', str(demo),
        '--today', '2026-10-01',
    ], env_file=None)
    assert code == 0
    assert target.exists()
    assert not demo.exists()
    products = {row[0] for row in _rows(target)}
    sold = pd.read_excel(workbook_path, sheet_name='Sales')
    recent = sold[(sold['Date'] > '2025-10-01') & (sold['Quantity'] > 0)]
    assert products == set(recent['Product'])
