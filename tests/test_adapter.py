import pandas as pd
import pytest

from src.adapters.boutique_xlsx import (
    SchemaError,
    attach_unit_costs,
    load_workbook_data,
)
from src.privacy import PrivacyError
from src.schema import STANDARD_COLUMNS

SALT = 'unit-test-salt-0123456789'


@pytest.fixture(scope='module')
def data(workbook_path):
    return load_workbook_data(workbook_path, salt=SALT)


@pytest.fixture(scope='module')
def raw_sales(workbook_path):
    return pd.read_excel(workbook_path, sheet_name='Sales')


def test_sales_use_the_standard_columns(data):
    assert list(data.sales.columns) == STANDARD_COLUMNS


def test_row_count_is_preserved(data, raw_sales):
    assert len(data.sales) == len(raw_sales)


def test_dates_quantities_and_prices_are_typed(data):
    sales = data.sales
    assert pd.api.types.is_datetime64_any_dtype(sales['date'])
    assert pd.api.types.is_integer_dtype(sales['quantity'])
    assert pd.api.types.is_float_dtype(sales['unit_price'])


def test_raw_customer_names_never_reach_the_output(data, raw_sales):
    names = set(raw_sales['Customer'].dropna())
    values = set(data.sales.astype(str).to_numpy().ravel())
    assert names.isdisjoint(values)


def test_unwanted_personal_columns_are_dropped(data):
    forbidden = {'customer', 'suburb', 'payment_method', 'Customer', 'Suburb'}
    assert forbidden.isdisjoint(data.sales.columns)


def test_customer_ids_are_hashed_and_stable(data, raw_sales):
    ids = data.sales['customer_id'].dropna()
    assert ids.str.fullmatch(r'[0-9a-f]{16}').all()
    named = raw_sales['Customer'].notna()
    assert data.sales.loc[named, 'customer_id'].notna().all()
    assert data.sales.loc[~named, 'customer_id'].isna().all()


def test_same_name_and_suburb_share_one_id(data, raw_sales):
    named = raw_sales[raw_sales['Customer'].notna()]
    expected = len(named[['Customer', 'Suburb']].drop_duplicates())
    assert data.sales['customer_id'].nunique() == expected


def test_refunds_are_linked_to_their_original_sale(data):
    sales = data.sales
    is_refund = sales['sale_id'].str.endswith('-R')
    assert is_refund.any()
    refunds = sales[is_refund]
    assert (refunds['refund_of'] == refunds['sale_id'].str[:-2]).all()
    assert sales.loc[~is_refund, 'refund_of'].isna().all()


def test_unit_cost_uses_the_buy_price_for_the_sale_year(data):
    prices = data.buy_prices.set_index(['product', 'year'])['unit_cost']
    row = data.sales.iloc[0]
    expected = prices[(row['product'], row['date'].year)]
    assert row['unit_cost'] == expected
    assert not row['cost_is_estimated']


def test_buy_prices_are_long_format(data):
    assert list(data.buy_prices.columns) == ['product', 'year', 'unit_cost']
    assert data.buy_prices['year'].between(2015, 2100).all()


def test_stock_sheet_is_parsed(data):
    assert list(data.stock.columns) == [
        'product', 'stock_on_hand', 'count_date', 'weeks_to_arrive',
        'min_order_qty', 'on_order',
    ]
    assert len(data.stock) > 0
    assert pd.api.types.is_datetime64_any_dtype(data.stock['count_date'])


def test_summary_figures_are_read_for_reconciliation(data, raw_sales):
    assert data.summary['number_of_refunds'] == (raw_sales['Quantity'] < 0).sum()
    assert data.summary['total_sales'] == pytest.approx(
        raw_sales['Sale Total (AUD)'].sum()
    )


def test_line_total_is_kept_for_validation(data):
    assert 'line_total' in data.sales.columns


def test_missing_salt_fails_closed(workbook_path, monkeypatch):
    monkeypatch.delenv('CUSTOMER_HASH_SALT', raising=False)
    with pytest.raises(PrivacyError):
        load_workbook_data(workbook_path)


def test_missing_sales_column_is_a_clear_error(workbook_path, tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    wb['Sales']['E1'] = 'Kind'  # renames the 'Category' header
    broken = tmp_path / 'broken.xlsx'
    wb.save(broken)
    with pytest.raises(SchemaError, match='Category'):
        load_workbook_data(broken, salt=SALT)


def test_missing_buy_prices_sheet_is_a_clear_error(workbook_path, tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    del wb['Buy Prices']
    broken = tmp_path / 'no_prices.xlsx'
    wb.save(broken)
    with pytest.raises(SchemaError, match='Buy Prices'):
        load_workbook_data(broken, salt=SALT)


def test_stock_sheet_is_optional(workbook_path, tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    del wb['Stock & Orders']
    path = tmp_path / 'no_stock.xlsx'
    wb.save(path)
    assert load_workbook_data(path, salt=SALT).stock is None


def test_unfilled_stock_rows_are_skipped_and_min_order_defaults_to_one(
        workbook_path, tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    ws = wb['Stock & Orders']
    for row in range(3, ws.max_row + 1):  # leave only the first product filled
        for col in range(2, 5):
            ws.cell(row=row, column=col).value = None
    ws['E2'] = None
    path = tmp_path / 'partial_stock.xlsx'
    wb.save(path)
    stock = load_workbook_data(path, salt=SALT).stock
    assert len(stock) == 1
    assert stock['min_order_qty'].iloc[0] == 1


def test_on_order_is_optional_and_defaults_to_zero(workbook_path, tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    wb['Stock & Orders'].delete_cols(6)  # remove the 'On Order' column
    path = tmp_path / 'no_on_order.xlsx'
    wb.save(path)
    stock = load_workbook_data(path, salt=SALT).stock
    assert (stock['on_order'] == 0).all()


def test_a_stock_sheet_with_nothing_filled_in_counts_as_missing(
        workbook_path, tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    ws = wb['Stock & Orders']
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=2).value = None
    path = tmp_path / 'empty_stock.xlsx'
    wb.save(path)
    assert load_workbook_data(path, salt=SALT).stock is None


def test_blank_totals_mean_formulas_were_not_saved(make_workbook):
    path = make_workbook(problems={'blank_total'})
    with pytest.raises(SchemaError, match='saved'):
        load_workbook_data(path, salt=SALT)


def test_unreadable_dates_are_rejected(workbook_path, tmp_path):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    wb['Sales']['A2'] = 'not a date'
    path = tmp_path / 'bad_date.xlsx'
    wb.save(path)
    with pytest.raises(SchemaError, match='Date'):
        load_workbook_data(path, salt=SALT)


def _sales_frame(rows):
    return pd.DataFrame(rows, columns=['product', 'date'])


def test_attach_costs_matches_product_and_year():
    prices = pd.DataFrame({
        'product': ['A', 'A', 'B'], 'year': [2024, 2025, 2025],
        'unit_cost': [10.0, 12.0, 5.0],
    })
    sales = _sales_frame([
        ('A', pd.Timestamp('2024-06-01')), ('A', pd.Timestamp('2025-06-01')),
        ('B', pd.Timestamp('2025-01-01')),
    ])
    result = attach_unit_costs(sales, prices)
    assert result['unit_cost'].tolist() == [10.0, 12.0, 5.0]
    assert not result['cost_is_estimated'].any()


def test_attach_costs_falls_back_to_latest_known_year_and_flags_it():
    prices = pd.DataFrame({
        'product': ['A', 'A'], 'year': [2024, 2025], 'unit_cost': [10.0, 12.0],
    })
    sales = _sales_frame([('A', pd.Timestamp('2027-03-01'))])
    result = attach_unit_costs(sales, prices)
    assert result['unit_cost'].iloc[0] == 12.0
    assert result['cost_is_estimated'].iloc[0]


def test_attach_costs_leaves_unknown_products_empty():
    prices = pd.DataFrame({'product': ['A'], 'year': [2025], 'unit_cost': [1.0]})
    sales = _sales_frame([('Mystery', pd.Timestamp('2025-03-01'))])
    result = attach_unit_costs(sales, prices)
    assert pd.isna(result['unit_cost'].iloc[0])


def _header_cell(ws, name, row=1):
    return next(c for c in ws[row] if c.value == name)


def _broken_copy(workbook_path, tmp_path, edit):
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path)
    edit(wb)
    path = tmp_path / 'edited.xlsx'
    wb.save(path)
    return path


def test_empty_quantity_is_a_clear_error(workbook_path, tmp_path):
    def edit(wb):
        ws = wb['Sales']
        ws.cell(row=2, column=_header_cell(ws, 'Quantity').column).value = None
    path = _broken_copy(workbook_path, tmp_path, edit)
    with pytest.raises(SchemaError, match='Quantity'):
        load_workbook_data(path, salt=SALT)


def test_text_in_a_price_cell_is_a_clear_error(workbook_path, tmp_path):
    def edit(wb):
        ws = wb['Sales']
        col = _header_cell(ws, 'Unit Price (AUD)').column
        ws.cell(row=2, column=col).value = 'free'
    path = _broken_copy(workbook_path, tmp_path, edit)
    with pytest.raises(SchemaError, match='Unit Price'):
        load_workbook_data(path, salt=SALT)


def test_text_in_a_buy_price_is_a_clear_error(workbook_path, tmp_path):
    def edit(wb):
        ws = wb['Buy Prices']
        header = next(r for r in ws.iter_rows() if r[0].value == 'Product')[0].row
        ws.cell(row=header + 1, column=3).value = 'tbc'
    path = _broken_copy(workbook_path, tmp_path, edit)
    with pytest.raises(SchemaError, match='Buy Prices'):
        load_workbook_data(path, salt=SALT)


def test_unreadable_stock_count_date_is_a_clear_error(workbook_path, tmp_path):
    def edit(wb):
        ws = wb['Stock & Orders']
        ws.cell(row=2, column=_header_cell(ws, 'Count Date').column).value = 'last week'
    path = _broken_copy(workbook_path, tmp_path, edit)
    with pytest.raises(SchemaError, match='Stock'):
        load_workbook_data(path, salt=SALT)
