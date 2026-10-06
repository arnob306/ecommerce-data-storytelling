"""
Synthetic boutique workbook generator.

Produces an .xlsx with the same sheets and column layout as the real sales
workbook, filled with invented data: made-up customer names, perturbed prices
and costs, and plausible festival seasonality. It is used for the public demo,
the tests and README screenshots, so real data never has to leave
data/private/.

Run:
    python -m src.synthetic.generator
Writes:
    data/synthetic/boutique_synthetic.xlsx
"""

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook

DEFAULT_OUTPUT = Path('data/synthetic/boutique_synthetic.xlsx')
DEFAULT_START = '2022-01-01'
DEFAULT_END = '2026-09-30'
BASE_YEAR = 2019

BASE_SALES_PER_DAY = 1.3
REFUND_RATE = 0.025
POOL_SIZE = 380
CUSTOMER_SKEW = 0.8
BLANK_NAME_RATE = 0.04
FESTIVAL_LABEL_RATE = 0.75
SAREE_FESTIVAL_BOOST = 2.0
WEEKDAY_MULT = {0: 0.8, 1: 0.8, 2: 0.9, 3: 0.95, 4: 1.1, 5: 1.5, 6: 1.3}

SALES_COLUMNS = [
    'Date', 'Sale ID', 'Customer', 'Product', 'Category', 'Quantity',
    'Unit Price (AUD)', 'Sale Total (AUD)', 'Payment Method', 'Sales Channel',
    'Customer Type', 'Suburb', 'Festival / Occasion', 'Business Event',
]
STOCK_COLUMNS = [
    'Product', 'Stock On Hand', 'Count Date', 'Weeks To Arrive',
    'Min Order Qty', 'On Order',
]
VALID_PROBLEMS = frozenset({
    'duplicate_id', 'orphan_refund', 'unknown_product', 'blank_total',
    'bad_total',
})


@dataclass(frozen=True)
class Product:
    name: str
    category: str
    base_price: float
    cost_ratio: float
    popularity: int


# Prices and cost ratios are invented, not copied from any real business.
PRODUCTS = (
    Product('Bindi Pack', 'Accessories', 3.0, 0.45, 8),
    Product('Potli Bag', 'Accessories', 19.0, 0.46, 3),
    Product('Embroidered Shawl', 'Accessories', 52.0, 0.50, 1),
    Product('Hair Clip Set', 'Accessories', 6.5, 0.48, 4),
    Product('Glass Bangles', 'Bangles', 8.5, 0.47, 12),
    Product('Gold-look Bangles', 'Bangles', 26.0, 0.50, 6),
    Product("Men's Cotton Panjabi", 'Clothing', 44.0, 0.58, 4),
    Product('Ladies Cotton Kurta', 'Clothing', 33.0, 0.57, 7),
    Product('Kids Panjabi Set', 'Clothing', 28.0, 0.56, 3),
    Product('Necklace Set', 'Jewellery', 38.0, 0.49, 5),
    Product('Jhumka Earrings', 'Jewellery', 16.0, 0.52, 8),
    Product('Cotton Saree', 'Saree', 42.0, 0.56, 6),
    Product('Jamdani Saree', 'Saree', 150.0, 0.57, 3),
    Product('Kantha Saree', 'Saree', 118.0, 0.60, 3),
    Product('Silk Saree', 'Saree', 215.0, 0.58, 2),
)

# (name, (start month, day), (end month, day), demand multiplier)
FESTIVALS = (
    ('Saraswati Puja', (1, 25), (2, 3), 2.0),
    ('Pohela Boishakh', (4, 8), (4, 16), 3.0),
    ('Eid', (4, 18), (4, 24), 1.6),
    ('Durga Puja', (10, 8), (10, 22), 3.0),
    ('Diwali / Kali Puja', (10, 28), (11, 10), 2.0),
    ('Christmas', (12, 12), (12, 24), 1.8),
)

# (event name, start, end, demand multiplier, product that cannot be sold)
EVENTS = (
    ('Silk Saree out of stock', '2024-08-05', '2024-09-15', 1.0, 'Silk Saree'),
    ('Price increase', '2023-02-01', '2023-02-14', 1.0, None),
    ('New stock from Bangladesh', '2023-06-05', '2023-06-18', 1.0, None),
    ('Viral Instagram post', '2025-05-12', '2025-05-25', 2.0, None),
)

# Deliberately invented names, so a demo customer can never match a real one.
FIRST_NAMES = (
    'Zorvan', 'Kelmira', 'Tavora', 'Nirelle', 'Quenby', 'Ostrel', 'Vandra',
    'Merrin', 'Jaskel', 'Ulvane', 'Pelvani', 'Wexel', 'Yarnel', 'Fenlow',
    'Gavrel', 'Xanvi', 'Lorkin', 'Mivrel', 'Tamsel', 'Orvane', 'Belkin',
    'Cavrel', 'Dunvar', 'Elvane', 'Ishvel', 'Norvik', 'Sarvane', 'Tolvem',
    'Uskara', 'Varnel',
)
SUBURBS = (
    'Dandenong', 'Noble Park', 'Springvale', 'Clayton', 'Glen Waverley',
    'Keysborough', 'Werribee', 'Other Melbourne',
)
PAYMENT_METHODS = ('Cash', 'Card', 'Bank Transfer', 'PayID')
CHANNELS = (
    'Facebook/Instagram', 'Word of Mouth', 'Home Pickup', 'Community Event',
)


def _price_for(product: Product, year: int) -> float:
    price = product.base_price * 1.045 ** (year - BASE_YEAR)
    return round(price * 2) / 2


def _cost_table(rng: np.random.Generator, years: list) -> dict:
    """Invented buy price per product per year."""
    table = {}
    for product in PRODUCTS:
        for year in years:
            noise = rng.uniform(0.95, 1.05)
            raw = product.base_price * product.cost_ratio
            cost = raw * 1.05 ** (year - BASE_YEAR) * noise
            table[(product.name, year)] = max(round(cost * 2) / 2, 0.5)
    return table


def _festival_on(day: pd.Timestamp) -> Optional[tuple]:
    for name, (sm, sd), (em, ed), mult in FESTIVALS:
        if (sm, sd) <= (day.month, day.day) <= (em, ed):
            return name, mult
    return None


def _event_on(day: pd.Timestamp) -> Optional[tuple]:
    for name, start, end, mult, blocked in EVENTS:
        if pd.Timestamp(start) <= day <= pd.Timestamp(end):
            return name, mult, blocked
    return None


def _build_customers(rng: np.random.Generator) -> tuple:
    """Return (customers, pick probabilities) with a few frequent buyers."""
    customers = []
    for _ in range(POOL_SIZE):
        first = FIRST_NAMES[rng.integers(len(FIRST_NAMES))]
        initial = chr(ord('A') + int(rng.integers(26)))
        customers.append({
            'name': f'{first} {initial}.',
            'suburb': SUBURBS[rng.integers(len(SUBURBS))],
        })
    weights = 1.0 / np.arange(1, POOL_SIZE + 1) ** CUSTOMER_SKEW
    return customers, weights / weights.sum()


def _pick_product(rng, boosted: bool, blocked: Optional[str]) -> Product:
    weights = np.array([
        p.popularity
        * (SAREE_FESTIVAL_BOOST if boosted and p.category == 'Saree' else 1)
        * (0 if p.name == blocked else 1)
        for p in PRODUCTS
    ], dtype=float)
    return PRODUCTS[rng.choice(len(PRODUCTS), p=weights / weights.sum())]


def _quantity(rng, product: Product) -> int:
    if product.base_price >= 30:
        return 1
    return int(min(rng.geometric(0.6), 4))


def _sales_for_day(rng, day, pool, seen, counter) -> list:
    customers, probs = pool
    festival = _festival_on(day)
    event = _event_on(day)
    mult = WEEKDAY_MULT[day.dayofweek]
    mult *= festival[1] if festival else 1.0
    mult *= event[1] if event else 1.0
    blocked = event[2] if event else None
    rows = []
    for _ in range(rng.poisson(BASE_SALES_PER_DAY * mult)):
        counter[0] += 1
        product = _pick_product(rng, festival is not None, blocked)
        customer = customers[rng.choice(len(customers), p=probs)]
        is_blank = rng.random() < BLANK_NAME_RATE
        key = (customer['name'], customer['suburb'])
        kind = 'Community Customer' if is_blank else (
            'Regular Customer' if key in seen else 'New Customer')
        seen.add(key)
        price = _price_for(product, day.year)
        qty = _quantity(rng, product)
        labelled = festival and rng.random() < FESTIVAL_LABEL_RATE
        rows.append({
            'Date': day.to_pydatetime(), 'Sale ID': f'S{counter[0]:05d}',
            'Customer': None if is_blank else customer['name'],
            'Product': product.name, 'Category': product.category,
            'Quantity': qty, 'Unit Price (AUD)': price,
            'Sale Total (AUD)': round(qty * price, 2),
            'Payment Method': PAYMENT_METHODS[rng.integers(len(PAYMENT_METHODS))],
            'Sales Channel': CHANNELS[rng.integers(len(CHANNELS))],
            'Customer Type': kind, 'Suburb': customer['suburb'],
            'Festival / Occasion': festival[0] if labelled else None,
            'Business Event': event[0] if event else None,
        })
    return rows


def _refund_for(rng, sale: dict, end: pd.Timestamp) -> Optional[dict]:
    if rng.random() >= REFUND_RATE:
        return None
    days = int(rng.integers(3, 20))
    when = pd.Timestamp(sale['Date']) + pd.Timedelta(days=days)
    if when > end:
        return None
    return {
        **sale, 'Date': when.to_pydatetime(),
        'Sale ID': f"{sale['Sale ID']}-R",
        'Quantity': -sale['Quantity'],
        'Sale Total (AUD)': -sale['Sale Total (AUD)'],
        'Festival / Occasion': None, 'Business Event': 'Return / refund',
    }


def _build_sales(rng, start, end) -> pd.DataFrame:
    pool = _build_customers(rng)
    seen: set = set()
    counter = [0]
    rows = []
    for day in pd.date_range(start, end, freq='D'):
        for sale in _sales_for_day(rng, day, pool, seen, counter):
            rows.append(sale)
            refund = _refund_for(rng, sale, pd.Timestamp(end))
            if refund:
                rows.append(refund)
    df = pd.DataFrame(rows)
    df = df.sort_values('Date', kind='stable').reset_index(drop=True)
    return df[SALES_COLUMNS]


def _inject_problems(df: pd.DataFrame, problems: set) -> pd.DataFrame:
    df = df.copy()
    positive = df.index[df['Quantity'] > 0]
    if 'duplicate_id' in problems:
        df.loc[positive[11], 'Sale ID'] = df.loc[positive[10], 'Sale ID']
    if 'unknown_product' in problems:
        df.loc[positive[20], 'Product'] = 'Mystery Item'
    if 'blank_total' in problems:
        df.loc[positive[30:33], 'Sale Total (AUD)'] = None
    if 'bad_total' in problems:
        df.loc[positive[40], 'Sale Total (AUD)'] += 5
    if 'orphan_refund' in problems:
        orphan = df.loc[positive[50]].copy()
        orphan['Sale ID'] = 'S99999-R'
        orphan['Quantity'] = -abs(orphan['Quantity'])
        orphan['Sale Total (AUD)'] = -abs(orphan['Sale Total (AUD)'])
        df = pd.concat([df, orphan.to_frame().T], ignore_index=True)
    return df


def _write_rows(ws, rows: Iterable, date_columns: tuple = ()) -> None:
    for row in rows:
        ws.append([None if pd.isna(v) else v for v in row])
    for idx in date_columns:
        for row in ws.iter_rows(min_row=2, min_col=idx + 1, max_col=idx + 1):
            cell = row[0]
            if isinstance(cell.value, datetime):
                cell.number_format = 'yyyy-mm-dd'


def _summary_rows(sales: pd.DataFrame, start: str, end: str) -> list:
    quantity = pd.to_numeric(sales['Quantity'])
    totals = pd.to_numeric(sales['Sale Total (AUD)'], errors='coerce')
    return [
        ['Sales Summary', None],
        [f'{start} to {end} - All figures in AUD (synthetic data)', None],
        [None, None],
        ['Key figures', None],
        ['Total sales (AUD, after refunds)', float(totals.sum())],
        ['Number of sales', int((quantity > 0).sum())],
        ['Number of returns / refunds', int((quantity < 0).sum())],
        ['Refunds (AUD)', float(-totals[quantity < 0].sum())],
    ]


def _buy_price_rows(costs: dict, years: list) -> list:
    rows = [
        ['What I paid for each item (AUD per item) - synthetic'],
        ['Rough buy price per year.'],
        [],
        ['Product', 'Category', *years],
    ]
    for product in PRODUCTS:
        rows.append([product.name, product.category,
                     *[costs[(product.name, y)] for y in years]])
    return rows + _profit_table_rows()


def _profit_table_rows() -> list:
    """The real sheet has a second 'Product' table below the prices."""
    rows = [
        [],
        ['Estimated profit by product (all years)'],
        ['Product', 'Items sold', 'Sales (AUD)', 'Cost of items (AUD)',
         'Profit (AUD)', 'Profit %'],
    ]
    rows += [[p.name, 0, 0.0, 0.0, 0.0, 0.0] for p in PRODUCTS]
    rows += [['Total', 0, 0.0, 0.0, 0.0, 0.0], [], ['Notes']]
    return rows


def _stock_rows(rng, end: str) -> list:
    rows = [STOCK_COLUMNS]
    for product in PRODUCTS:
        rows.append([
            product.name, int(rng.integers(0, 14)),
            pd.Timestamp(end).to_pydatetime(), int(rng.integers(3, 9)),
            int(rng.choice([1, 3, 5, 10])), int(rng.choice([0, 0, 0, 3, 5])),
        ])
    return rows


def generate_workbook(
    path,
    seed: int = 7,
    start: str = DEFAULT_START,
    end: str = DEFAULT_END,
    problems: Optional[set] = None,
) -> Path:
    """Write a synthetic workbook to ``path`` and return the path."""
    problems = set(problems or ())
    unknown = problems - VALID_PROBLEMS
    if unknown:
        raise ValueError(f'Unknown problems: {sorted(unknown)}')

    rng = np.random.default_rng(seed)
    years = list(range(BASE_YEAR, pd.Timestamp(end).year + 1))
    costs = _cost_table(rng, years)
    sales = _inject_problems(_build_sales(rng, start, end), problems)

    wb = Workbook()
    ws_sales = wb.active
    ws_sales.title = 'Sales'
    ws_sales.append(SALES_COLUMNS)
    _write_rows(ws_sales, sales[SALES_COLUMNS].itertuples(index=False), (0,))
    _write_rows(wb.create_sheet('Simple Summary'),
                _summary_rows(sales, start, end))
    _write_rows(wb.create_sheet('Buy Prices'), _buy_price_rows(costs, years))
    _write_rows(wb.create_sheet('Stock & Orders'), _stock_rows(rng, end), (2,))

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == '__main__':
    out = generate_workbook(DEFAULT_OUTPUT)
    print(f'Wrote {out} ({out.stat().st_size / 1024:.0f} KB)')
